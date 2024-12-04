import nltk
import time
import os
import openai

# nltk.download("punkt")
from langchain_community.document_loaders import (
    UnstructuredWordDocumentLoader,
    UnstructuredPDFLoader,
)
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.vectorstores import Chroma

from langchain_community.chat_models import ChatOllama

from langchain.chains import RetrievalQA
from langchain_openai import ChatOpenAI, OpenAIEmbeddings


# Library for Reading Template file
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

# Library for Cost Calculations
import tiktoken

from dotenv import load_dotenv

load_dotenv()

openai.api_key = os.environ["OPENAI_API_KEY"]


def num_tokens_from_string(string: str, encoding_name: str) -> int:
    """Returns the number of tokens in a text string."""
    encoding = tiktoken.get_encoding(encoding_name)
    num_tokens = len(encoding.encode(string))
    return num_tokens


def calculate_embedding_cost(num_tokens):
    """Calculate the cost of embeddings based on current OpenAI pricing."""
    # Price per 1K tokens for text-embedding-ada-002, as of September 2023
    price_per_1k_tokens = 0.0001
    return (num_tokens / 1000) * price_per_1k_tokens


def docx_to_markdown(docx_path):
    doc = docx.Document(docx_path)
    markdown_content = []

    def convert_table_to_markdown(table):
        markdown_table = []
        max_col_widths = [0] * len(table.rows[0].cells)

        # First pass: determine maximum width for each column
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                max_col_widths[i] = max(max_col_widths[i], len(cell.text.strip()))

        # Second pass: create markdown table
        for i, row in enumerate(table.rows):
            markdown_row = []
            for j, cell in enumerate(row.cells):
                cell_content = cell.text.strip().ljust(max_col_widths[j])
                markdown_row.append(cell_content)
            markdown_table.append("| " + " | ".join(markdown_row) + " |")

            # Add header separator after first row
            if i == 0:
                separator = (
                    "|"
                    + "|".join(["-" * (width + 2) for width in max_col_widths])
                    + "|"
                )
                markdown_table.append(separator)

        return "\n".join(markdown_table)

    for element in doc.element.body:
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, doc)
            if paragraph.text.strip():  # Only add non-empty paragraphs
                markdown_content.append(paragraph.text)
        elif isinstance(element, CT_Tbl):
            table = Table(element, doc)
            markdown_content.append(convert_table_to_markdown(table))
            markdown_content.append("")  # Add an empty line after the table

    # Handle text boxes (shapes)
    for shape in doc.inline_shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.text.strip():  # Only add non-empty paragraphs
                    markdown_content.append(f"[Text Box] {paragraph.text}")
    return "\n\n".join(markdown_content)


def load_docx(file_path):
    loader = UnstructuredWordDocumentLoader(file_path)
    documents = loader.load()
    print(f"Loaded {len(documents)} documents")
    return documents


def split_documents(documents):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=5000, chunk_overlap=3500)
    chunks = text_splitter.split_documents(documents)
    print(f"Split into {len(chunks)} chunks")
    return chunks, len(chunks)


def create_vector_store(chunks):
    total_tokens = 0
    for chunk in chunks:
        total_tokens += num_tokens_from_string(chunk.page_content, "cl100k_base")

    estimated_cost = calculate_embedding_cost(total_tokens)

    print(f"Total tokens to be embedded: {total_tokens}")
    print(f"Estimated cost for embeddings: ${estimated_cost:.4f}")
    vector_db = Chroma.from_documents(
        documents=chunks,
        embedding=OpenAIEmbeddings(model="text-embedding-ada-002"),
        collection_name="local-project-rag",
    )
    return vector_db


def setup_llm():
    # local_model = "llama3.1:latest"
    llm = ChatOpenAI(model="gpt-4o", temperature=0, max_retries=2)
    return llm


def setup_activity_llm():
    llm = ChatOpenAI(model="gpt-4o", temperature=0.2)
    return llm


def project_scenario_extraction_pipeline(vector_store, llm):
    query = """
    You are a professional document analyzer. As a owner of this document, I am asking you to follow my guideline.
    Based on the content of the document, please extract only the scenario described in Assessment Task 2 for helping a student. Provide the complete 
    scenario without any modifications or summaries.

    Please provide the output in the following format:
    Scenario: [Extracted scenario from Assessment Task 2]
    """

    retriever = vector_store.as_retriever()
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, chain_type="stuff", retriever=retriever, return_source_documents=True
    )

    result = qa_chain.invoke({"query": query})

    return result["result"]


def project_activities_extraction_pipeline(vector_store, llm, scenario, chunk_size):
    query = """
    You are a professional document analyzer. Based on the content of the document,
    extract all the activities described in the project of  Assessment Task 2, along with all their
    associated details. For your help, I marked the starting point of the activities
    like "Activities-" and for each activity, I have marked the requirements also like  "Requirements:"
    They are stated after the {scenario}.

    Please provide the output in the following format:
    Activities:
        1. [Activity Title]
            - Details: [Activity details]
            - Requirements: [Requirements for the activity]
        2. [Activity Title]
            - Details: [Activity details]
            - Requirements: [Requirements for the activity]
        continue for all activities

    Ensure that you capture all the information provided for each activity without
    any modifications or summaries.
    """
    retriever = vector_store.as_retriever(search_kwargs={"k": min(chunk_size, 10)})
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, chain_type="stuff", retriever=retriever, return_source_documents=True
    )

    result = qa_chain.invoke({"query": query})

    return result["result"]


def generate_project_output(llm, project_info, template_content):
    prompt = f"""
    [INST] You are a creative scenario generator for project-based learning.Your task is to create or assume a detailed, elaborated, engaging scenario based on the provided project information and only include the scenario details. After generating the scenario, you have to complete the tasks by strictly following the template structure. 
    Project Information: {project_info}
    Template Structure: {template_content}
    [/INST]
    """
    response = llm.invoke(prompt)
    return response.content


def generate_project_scenario(llm, project_info):

    query = f"""
    You're a professional and creative scenario generator for project-based learning. Your task is to create or assume a engaging scenario based on the provided project information. 
    
    Project Information: {project_info}
    Output format:
    Generated Scenario: [Scenario Summary]
    """
    response = llm.invoke(query)
    return response.content


def extract_activities(activities_text):
    """Extract activities into a structured list"""
    activities = []
    current_activity = None

    # Split by lines and clean up the text
    lines = activities_text.split("\n")

    for line in lines:
        line = line.strip()

        # Skip empty lines and the "Activities:" header
        if not line or line == "Activities:":
            continue

        # Check for new activity (starts with number followed by period)
        if line[0].isdigit() and ". " in line:
            if current_activity:
                activities.append(current_activity)

            activity_number = int(line.split(".")[0])
            activity_title = line.split(".")[1].strip()
            current_activity = {
                "number": activity_number,
                "title": activity_title,
                "details": "",
                "requirements": "",
            }

        # Check for details and requirements
        elif "- Details:" in line:
            current_activity["details"] = line.split("- Details:")[1].strip()
        elif "- Requirements:" in line:
            current_activity["requirements"] = line.split("- Requirements:")[1].strip()

    # Don't forget to append the last activity
    if current_activity:
        activities.append(current_activity)

    return activities

def extract_guideline_per_activity(guideline_vectorstore,llm, activity, scenario, template):
    query = f"""[INST] 
    You are a professional guideline extraction chatbot. 
    Based on the provided context, first analyze the whole context of the document and the provided scenario : {scenario} and then given template : {template}.
    Then extract the guidelines related the given activity. Given :
        Activity Title: {activity['title']}
        Details: {activity['details']}
        Requirements: {activity['requirements']}
    If a template should be used like activity details can be said that "template provided", format the response accordingly like the {template}.
    Otherwise provide a freeform response.

    If guideline does not exist for the given activity return guidline not available.
    
    [/INST]
    """


    retriever = guideline_vectorstore.as_retriever()
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, chain_type="stuff", retriever=retriever,
    )

    result = qa_chain.invoke({"query": query})
    return result["result"]
def extract_notes_per_activity(note_vectorstore,llm, activity,scenario,template):
    query = f"""[INST] 
    You are a professional information extraction chatbot. 
    Based on the provided context, first analyze the whole context of the document and the provided scenario : {scenario} and then given template : {template}.
    Then extract the important information related the given activity. Given :
        Activity Title: {activity['title']}
        Details: {activity['details']}
        Requirements: {activity['requirements']}
    If a template should be used like activity details can be said that "template provided", format the response accordingly like the {template}.
    Otherwise provide a freeform response.

    If information does not exist for the given activity return no related information available.
    
    [/INST]
    """

    retriever = note_vectorstore.as_retriever()
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, chain_type="stuff", retriever=retriever,
    )

    result = qa_chain.invoke({"query": query})
    return result["result"]

def process_single_activity(llm, activity, template, scenario, guideline, summarized_prev_activities = None, notes=None):

    prompt = f"""
    You're an intelligent assistant and you are smart enough to analyze and answer any task using the context given to you.
    Given the {scenario} [Strictly follow that scenario, don't modify it], the following activity needs to be completed:

    Activity Title: {activity['title']}
    Details: {activity['details']}
    Requirements: {activity['requirements']}

    Use the provided guideline : {guideline} to improve your response. 
    Analyze the previous activities : {summarized_prev_activities} to keep your response consistent with the previous activities.

    Use the provided notes : {notes} to improve your response if available

    Analyze the activity details and requirements  to determine if a template structure should be used for the response. If a template should be used like activity details can 
    be said that "template provided", format the response accordingly like the {template}. Otherwise, provide a freeform response like below:
    
    [Activity Number]: [Activity Title]
    [Answer]: [Generated Activity Answer]
    """

    response = llm.invoke(prompt)
    return response.content

def summarize_activities(activities_text, llm):
    activities_text_str = "\n".join(activities_text)
    prompt = f"""
    [INST] You are a professional document summarizer, Based on the given activities texts, summarize the activities.
    Activities:
    {activities_text_str}

    Remember to keep the critical information of each activity in the summary.
    [/INST]
    """

    response = llm.invoke(prompt)
    return response.content
def process_note_files(note_files):
    """Process multiple note files and create a combined vector store"""
    all_note_chunks = []
    
    for note_file in note_files:
        note_chunks, _ = split_documents(load_docx(note_file))
        all_note_chunks.extend(note_chunks)
        
    note_vector_db = create_vector_store(all_note_chunks)
    return note_vector_db

def process_guideline_files(guideline_files):
    """Process multiple guideline files and create a combined vector store"""
    all_guideline_chunks = []
    
    for guideline_file in guideline_files:
        guideline_chunks, _ = split_documents(load_docx(guideline_file))
        all_guideline_chunks.extend(guideline_chunks)
        
    guideline_vector_db = create_vector_store(all_guideline_chunks)
    return guideline_vector_db    
def main():
    project_tasks_file = "docs/SITXWHS006/SITXWHS006 Student Assessment Tasks 2 (Project) Short.docx"
    # relevant_docs_file = "../data/project/SITXWHS006 WHS Plan (Notes).docx"
    template_file = "docs/SITXWHS006/SITXWHS006 WHS Plan (Template).docx"
    documents = load_docx(project_tasks_file)
    chunks, chunk_size = split_documents(documents)
    vector_db = create_vector_store(chunks)
    guideline_files = ["docs/SITXWHS006/SITXWHS006 WHS Plan Assessor Guidelines.docx"]
    note_files = [
        "docs/SITXWHS006/SITXWHS006 WHS Plan (Notes).docx"
        # Add additional note files here
    ]
    note_vector_db = process_note_files(note_files)
    guideline_vector_db = process_guideline_files(guideline_files)
    llm = setup_llm()
    start_time = time.time()
    project_info = project_scenario_extraction_pipeline(vector_db, llm)
    project_scenario = generate_project_scenario(llm, project_info)
    project_activities = project_activities_extraction_pipeline(
        vector_db, llm, project_info, chunk_size
    )
    markdown_content = docx_to_markdown(template_file)
    print(
        f"==================================\n{project_info}\n=================================="
    )
    print(
        f"==================================\n{project_scenario}\n=================================="
    )

    activities = extract_activities(project_activities)
    print("*" * 40)
    print("Extracted activities:")
    for activity in activities:
        print(f"\nActivity {activity['number']}:")
        print(f"Title: {activity['title']}")
        print(f"Details: {activity['details']}")
        print(f"Requirements: {activity['requirements']}")
        print("-" * 40)

    activity_llm = setup_activity_llm()
    prev_activities = []
    summarize_activities_text = None
    for activity in activities:
        print(f"\nProcessing Activity {activity['number']}...")
        extracted_guideline = extract_guideline_per_activity(
            guideline_vector_db, activity_llm, activity, project_scenario, markdown_content
        )
        print("-" * 40)
        print(f"Extracted Guideline: {extracted_guideline}")
        print("-" * 40)

        extracted_notes = extract_notes_per_activity(
            note_vector_db, activity_llm, activity, project_scenario, markdown_content
        )
        print(f"Extracted Notes: {extracted_notes}")
        print("-" * 40)
        # Generate content for this activity
        if len(prev_activities) > 0:
            summarize_activities_text = summarize_activities(prev_activities, activity_llm)
        output_content = process_single_activity(
            activity_llm, activity, markdown_content, project_scenario, extracted_guideline, summarize_activities_text, extracted_notes
        )
        prev_activities.append(output_content)
        

        # Save to individual file
        filename = f"output/project/output_activity_{activity['number']}.md"
        with open(filename, "w") as f:
            f.write(f"# Activity {activity['number']}: {activity['title']}\n\n")
            f.write(output_content)

        print(f"Saved output to {filename}")

    end_time = time.time()
    print(f"Time spent: {(end_time-start_time)/60} minutes")


if __name__ == "__main__":
    main()
