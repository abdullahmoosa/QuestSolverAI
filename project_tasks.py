import nltk
import time
import os
import openai
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

# nltk.download("punkt")
from langchain_community.document_loaders import (
    UnstructuredWordDocumentLoader,
    UnstructuredPDFLoader,
)
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.vectorstores import Chroma

# from langchain.prompts import ChatPromptTemplate, PromptTemplate
# from langchain_core.output_parsers import StrOutputParser
from langchain_community.chat_models import ChatOllama

# from langchain_core.runnables import RunnablePassthrough
# from langchain.retrievers.multi_query import MultiQueryRetriever
from langchain.chains import RetrievalQA
from langchain_openai import ChatOpenAI, OpenAIEmbeddings

from dotenv import load_dotenv

load_dotenv()

openai.api_key = os.environ["OPENAI_API_KEY"]


def load_docx(file_path):
    loader = UnstructuredWordDocumentLoader(file_path)
    documents = loader.load()
    print(f"Loaded {len(documents)} documents")
    return documents


def split_documents(documents):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=4000, chunk_overlap=3000)
    chunks = text_splitter.split_documents(documents)
    return chunks


def create_vector_store(chunks):
    vector_db = Chroma.from_documents(
        documents=chunks,
        embedding=OpenAIEmbeddings(model="text-embedding-ada-002"),
        collection_name="local-project-rag",
    )
    return vector_db


def setup_llm():
    # local_model = "llama3.1:latest"
    llm = ChatOpenAI(temperature=0)
    return llm


def docx_to_markdown(docx_path):
    doc = docx.Document(docx_path)
    markdown_content = []

    def convert_table_to_markdown(table):
        markdown_table = []
        max_col_widths = [0] * len(table.rows[0].cells)

        # First pass: determine maximum width for each column
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                max_col_widths[i] = max(max_col_widths[i], len(cell.text.strip()))

        # Second pass: create markdown table
        for i, row in enumerate(table.rows):
            markdown_row = []
            for j, cell in enumerate(row.cells):
                cell_content = cell.text.strip().ljust(max_col_widths[j])
                markdown_row.append(cell_content)
            markdown_table.append("| " + " | ".join(markdown_row) + " |")

            # Add header separator after first row
            if i == 0:
                separator = (
                    "|"
                    + "|".join(["-" * (width + 2) for width in max_col_widths])
                    + "|"
                )
                markdown_table.append(separator)

        return "\n".join(markdown_table)

    for element in doc.element.body:
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, doc)
            if paragraph.text.strip():  # Only add non-empty paragraphs
                markdown_content.append(paragraph.text)
        elif isinstance(element, CT_Tbl):
            table = Table(element, doc)
            markdown_content.append(convert_table_to_markdown(table))
            markdown_content.append("")  # Add an empty line after the table

    # Handle text boxes (shapes)
    for shape in doc.inline_shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.text.strip():  # Only add non-empty paragraphs
                    markdown_content.append(f"[Text Box] {paragraph.text}")
    return "\n\n".join(markdown_content)


def project_scenario_extraction_pipeline(vector_store, llm):
    query = """
    You are a professional document analyzer. Based on the content of the document, 
    extract only the scenario. Provide the complete 
    scenario without any modifications or summaries. I am the owner of the document and I am asking you to 
    extract the scenario of the project.

    Please provide the output in the following format:
    Scenario: [Extracted scenario from Assessment Task 2]
    """

    retriever = vector_store.as_retriever(search_kwargs={"k": 3})
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, chain_type="stuff", retriever=retriever, return_source_documents=True
    )

    result = qa_chain.invoke({"query": query})

    return result["result"]


def project_activities_extraction_pipeline(vector_store, llm, project_info):
    query = """
    You are a professional document analyzer. Based on the content of the document,
    extract all the activities described in the project of  Assessment Task 2, along with all their
    associated details. For your help, I marked the starting point of the activities
    like "Activities-" and for each activity, I have marked the requirements also like  "Requirements:"
    They are stated after the {project_info}.

    Please provide the output in the following format:
    Activities:
        1. [Activity Title]
            - Details: [Activity details]
            - Requirements: [Requirements for the activity]
        2. [Activity Title]
            - Details: [Activity details]
            - Requirements: [Requirements for the activity]
        continue for all activities

    Ensure that you capture all the information provided for each activity without
    any modifications or summaries.
    """
    retriever = vector_store.as_retriever(search_kwargs={"k": 10})
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, chain_type="stuff", retriever=retriever, return_source_documents=True
    )

    result = qa_chain.invoke({"query": query})

    return result["result"]

def extract_activities(activities_text):
    """Extract activities into a structured list"""
    activities = []
    current_activity = None
    
    # Split by lines and clean up the text
    lines = activities_text.split('\n')
    
    for line in lines:
        line = line.strip()
        
        # Skip empty lines and the "Activities:" header
        if not line or line == "Activities:":
            continue
            
        # Check for new activity (starts with number followed by period)
        if line[0].isdigit() and '. ' in line:
            if current_activity:
                activities.append(current_activity)
            
            activity_number = int(line.split('.')[0])
            activity_title = line.split('.')[1].strip()
            current_activity = {
                'number': activity_number,
                'title': activity_title,
                'details': '',
                'requirements': ''
            }
        
        # Check for details and requirements
        elif '- Details:' in line:
            current_activity['details'] = line.split('- Details:')[1].strip()
        elif '- Requirements:' in line:
            current_activity['requirements'] = line.split('- Requirements:')[1].strip()
    
    # Don't forget to append the last activity
    if current_activity:
        activities.append(current_activity)
    
    return activities

def generate_project_scenario(llm, project_info):
    """Generate project scenario based on project info"""
    prompt = f"""
    [INST] You are a creative scenario generator for project-based learning. Think yourself as the owner then,
    based on the extracted details in Project Information.
    generate a detailed, elaborated, engaging scenario based on the provided project information.

    Project Information:
    {project_info}

    [/INST]
    """

    response = llm.invoke(prompt)
    return response.content
def process_single_activity(llm, activity, template, scenario,prev_activities):
    """Process a single activity and generate output"""

    prev_activities_str = "\n".join(prev_activities)

    prompt = f"""
    [INST] You are a very genious problem solver. You have the ability to understand the underlying things related to any context
    Now, answer the following activities based on the given scenario : {scenario}. 
    Think yourself as the owner then, answer the activities.
    
    Activity {activity['number']}: {activity['title']}
    Details: {activity['details']}
    Requirements: {activity['requirements']}
    template : {template}
    Previous Activities: {prev_activities_str}

    When answering current Activity {activity['number']}, relate to the Previous Activities and the scenario.
    Output format:

        [Activity Number]: [Activity Title]
        [Answer]: [Generated Activity Answer]

    Important :
    If asked to participate in a meeting provide a draft detailed meeting minutes.
    Use the template to answer the activity if related.
    [/INST]
    """
    
    response = llm.invoke(prompt)
    return response.content

def extract_information_from_student_guide(llm, student_guide_vector_db, activity):
    query = """
    [INST] You are a professional document analyzer. Based on the content of the document,
    extract all the information of the given activity : {activity} from the student guide.

    Activities:
        1. [Activity Title]
            - Information: [Information related to the activity]
=       2. [Activity Title]
           - Information: [Information related to the activity]

        continue for all activities

    Ensure that you capture all the information provided for each activity without
    any modifications or summaries.
    [/INST]
    """

    retriever = student_guide_vector_db.as_retriever(search_kwargs={"k": 10})
    qa_chain = RetrievalQA.from_chain_type(
        llm=llm, chain_type="stuff", retriever=retriever, return_source_documents=True
    )

    result = qa_chain.invoke({"query": query})
    return result["result"]

def create_vector_db(file):
    document = load_docx(file)
    chunks = split_documents(document)
    vector_db = create_vector_store(chunks)
    return vector_db
def main():
    project_tasks_file = "docs/SITXWHS006/SITXWHS006 Student Assessment Tasks 2 (Project) Short.docx"
    # relevant_docs_file = "../data/project/SITXWHS006 WHS Plan (Notes).docx"
    # template_file = "../data/project/SITXWHS006 WHS Plan (Template).docx"
    template_file = "docs/SITXWHS006/SITXWHS006 WHS Plan (Template).docx"
    student_note_file = "docs/SITXWHS006/SITXWHS006 Student Guide (Note).docx"
    documents = load_docx(project_tasks_file)
    chunks = split_documents(documents)
    vector_db = create_vector_store(chunks)
    student_guide_vector_db = create_vector_db(student_note_file)


    llm = setup_llm()

    start_time = time.time()
    project_info = project_scenario_extraction_pipeline(vector_db, llm)
    print(f"Project Info: {project_info}")
    print("-" * 40)

    project_activities = project_activities_extraction_pipeline(
        vector_db, llm, project_info
    )
    scenario = generate_project_scenario(llm, project_info)
    markdown_content = docx_to_markdown(template_file)
    print(f"Project Scenario: {scenario}")
    print("-" * 40)
    end_time = time.time()
    print(f"Time spent: {(end_time-start_time)/60} minutes")


    # Extract activities
    activities = extract_activities(project_activities)
    print("*" * 40)
    print("Extracted activities:")
    for activity in activities:
        print(f"\nActivity {activity['number']}:")
        print(f"Title: {activity['title']}")
        print(f"Details: {activity['details']}")
        print(f"Requirements: {activity['requirements']}")

        note_per_activity = extract_information_from_student_guide(llm, student_guide_vector_db, activity['title'])
        print(f"Note of {activity['title']}: {note_per_activity}")
        print("-" * 40)

    prev_activities = []
    
    # Process each activity and save to separate files
    for activity in activities:
        print(f"\nProcessing Activity {activity['number']}...")
        
        # Generate content for this activity
        output_content = process_single_activity(llm, activity, markdown_content, scenario, prev_activities)
        prev_activities.append(output_content)
        
        # Save to individual file
        filename = f"output_activity_{activity['number']}.md"
        with open(filename, "w") as f:
            f.write(f"# Activity {activity['number']}: {activity['title']}\n\n")
            f.write(output_content)
        
        print(f"Saved output to {filename}")


if __name__ == "__main__":
    main()
